"""The target-proposal deck renderer: `Slide` objects in, PowerPoint bytes out.

The slides are built by hand here, so nothing about a Report is needed: the
renderer only draws what it is given. What is checked is what a scientist would
notice - nothing hangs off the slide, long text continues on a "(cont.)" slide
instead of shrinking or running into the footer, each page's footer lists only
its own references, a figure goes where its shape says it should and is always
credited, and the marks that say "assembled by rule" and "integrity warning" are
there.
"""

from __future__ import annotations

import io
import subprocess
import sys

import pytest

pytest.importorskip("pptx")  # an optional extra: the module is skipped, not failed, without it

from PIL import Image  # noqa: E402  (python-pptx depends on Pillow)
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from bedrock.litreview.report.deck_model import (  # noqa: E402
    DECK_VERSION,
    FigureAsset,
    Para,
    RefMark,
    Slide,
    Visual,
    VisualItem,
)
from bedrock.litreview.report.deck_render import render  # noqa: E402

REFS = [
    RefMark(1, "Smith 2020", "https://example.org/smith"),
    RefMark(2, "Lee 2019", "https://example.org/lee"),
    RefMark(3, "Okafor 2021", None),
]
SENTENCE = (
    "IL-23 blockade reduced psoriasis area and severity index scores by 75 percent in the treated group "
    "compared with placebo over sixteen weeks in a randomised trial of adults with moderate to severe disease [{n}]."
)


def para(n: int, **kw) -> Para:
    return Para(SENTENCE.format(n=n), refs=[n], **kw)


def open_deck(slides: list[Slide]):
    return Presentation(io.BytesIO(render(slides)))


def texts(slide) -> list[str]:
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame]


def titles(prs) -> list[str]:
    return [next(s for s in slide.shapes if s.name == "Title").text_frame.text for slide in prs.slides if
            any(s.name == "Title" for s in slide.shapes)]


def named(slide, name: str):
    return next(s for s in slide.shapes if s.name == name)


def png(path, width: int, height: int):
    Image.new("RGB", (width, height), (200, 220, 240)).save(path)
    return str(path)


def figure(path: str, label: str = "Figure 2", **kw) -> FigureAsset:
    return FigureAsset(
        "p1", "PMC1", "F2", label, kw.pop("caption", "Serum IL-23 fell after treatment."), "f.png",
        kw.pop("licence", "CC BY 4.0"), kw.pop("licence_url", "https://creativecommons.org/licenses/by/4.0/"),
        ["E1"], path,
    )


def every_layout() -> list[Slide]:
    stats = Visual("stats_callout", [VisualItem("Papers read", value=14), VisualItem("Human studies", value=5)], "Counts")
    bars = Visual("bar_chart", [VisualItem("improved", value=9), VisualItem("no_effect", value=3), VisualItem("worsened", value=1)], "Findings")
    ring = Visual("doughnut_chart", [VisualItem("mouse", value=6), VisualItem("rat", value=2)])
    compare = Visual("comparison", [VisualItem("A alone", "Improved findings", 6, "positive"), VisualItem("Together", "None", 0, "negative")])
    stack = Visual("stacked_list", [VisualItem("IL23A", "12 papers"), VisualItem("TNF", "7 papers")])
    cards = Visual("callout_cards", [VisualItem("Psoriasis", "8 human"), VisualItem("Crohn's", "2 human")])
    slides = [Slide("IL-23 x TNF", "title", subtitle="Ayush Pawar | 30 Sep 2026\nWhat should I target?", refs=REFS)]
    slides += [Slide(f"Visual {v.type}", "content", [para(1)], visual=v, refs=REFS) for v in (stats, bars, ring, compare, stack, cards)]
    slides += [Slide("Text only", "content", [para(2), Para("A sub point", level=1)], refs=REFS)]
    slides += [Slide("Competitors", "table", table=[["Drug", "Ref."], ["Agent A", "[1]"], ["Agent B", "[2]"]], refs=REFS)]
    slides += [Slide("SWOT", "matrix", [Para("Strengths", bold=True), para(1), Para("Weaknesses", bold=True), para(2),
                                         Para("Opportunities", bold=True), Para("Threats", bold=True)], refs=REFS)]
    slides += [Slide("References", "references", [Para("[1] Smith 2020. https://example.org/smith", link="https://example.org/smith")], size=10)]
    return slides


def test_every_layout_and_visual_renders_inside_the_slide():
    slides = every_layout()
    prs = open_deck(slides)
    assert len(prs.slides) >= len(slides)
    width, height = prs.slide_width, prs.slide_height
    assert (width, height) == (Emu(int(13.333 * 914400)), Emu(int(7.5 * 914400)))
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (number, shape.name)
            assert shape.left + shape.width <= width, (number, shape.name)
            assert shape.top + shape.height <= height, (number, shape.name)


def test_charts_are_native_and_carry_the_counted_values():
    prs = open_deck(every_layout())
    charts = [s.chart for slide in prs.slides for s in slide.shapes if s.has_chart]
    assert len(charts) == 2
    bar, ring = charts
    assert list(bar.plots[0].categories) == ["improved", "no_effect", "worsened"]
    assert list(bar.plots[0].series[0].values) == [9.0, 3.0, 1.0]
    assert list(ring.plots[0].series[0].values) == [6.0, 2.0]
    # counts sit in the legend, so no number lies on a slice
    assert list(ring.plots[0].categories) == ["mouse (6)", "rat (2)"]


def test_a_bar_keeps_its_outcome_colour_whatever_its_position():
    bars = Visual("bar_chart", [VisualItem("Worsened", "worsened", 1, "negative"), VisualItem("No effect", "no_effect", 3, "neutral"),
                                VisualItem("improved", value=9)])
    prs = open_deck([Slide("Human", "content", [para(1)], visual=bars, refs=REFS)])
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    colours = [str(chart.plots[0].series[0].points[i].format.fill.fore_color.rgb) for i in (0, 1, 2)]
    assert colours == ["DC2626", "94A3B8", "42BA97"]  # by the outcome's key or its name, never by position
    assert list(chart.plots[0].categories) == ["Worsened", "No effect", "improved"]


def test_a_long_text_slide_continues_instead_of_shrinking_or_overflowing():
    long = Slide("Mechanism", "content", [para(1 + i % 3) for i in range(24)], refs=REFS, kicker="Animal and cell evidence")
    prs = open_deck([long])
    assert len(prs.slides) > 1
    assert titles(prs) == ["Mechanism"] + ["Mechanism (cont.)"] * (len(prs.slides) - 1)
    kept = []
    for slide in prs.slides:
        footer_top = min(s.top for s in slide.shapes if s.name.startswith("Footer"))
        body = [s for s in slide.shapes if s.name == "Body"]
        assert body and all(s.top + s.height <= footer_top for s in body)
        for box in body:
            for p in box.text_frame.paragraphs:
                kept += [p.text]
                assert {r.font.size.pt for r in p.runs} == {16}  # never shrunk
    assert kept == [SENTENCE.format(n=1 + i % 3) for i in range(24)]  # every sentence, once, unchanged


def test_a_heading_is_never_left_alone_at_the_bottom_of_a_page():
    # Slide the heading through every position near a page break: it must always
    # move to the next page together with the paragraph it introduces.
    starts_a_page = False
    for lead in range(1, 14):
        paras = [para(1) for _ in range(lead)] + [Para("Key limitations", bold=True), para(2, level=1), para(3, level=1)]
        prs = open_deck([Slide("Executive summary", "content", paras, refs=REFS)])
        for s in prs.slides:
            body = named(s, "Body").text_frame.paragraphs
            assert body[-1].text != "Key limitations", lead
            starts_a_page = starts_a_page or (s is not prs.slides[0] and body[0].text == "Key limitations")
    assert starts_a_page  # the loop did reach a page break at the heading


def test_sub_bullets_are_measured_so_they_cannot_run_into_the_footer():
    slide = Slide("Limits", "content", [Para(SENTENCE.format(n=1), level=1, refs=[1]) for _ in range(20)], refs=REFS)
    prs = open_deck([slide])
    assert len(prs.slides) > 1
    for s in prs.slides:
        footer_top = min(shape.top for shape in s.shapes if shape.name.startswith("Footer"))
        assert all(x.top + x.height <= footer_top for x in s.shapes if x.name == "Body")


def test_one_paragraph_taller_than_a_page_is_split_between_words_not_cut():
    words = " ".join(f"word{i}" for i in range(2500))
    prs = open_deck([Slide("Huge", "content", [Para(words)])])
    assert len(prs.slides) > 1
    rebuilt = " ".join(
        p.text for slide in prs.slides for s in slide.shapes if s.name == "Body" for p in s.text_frame.paragraphs
    )
    assert rebuilt == words


def test_footer_lists_only_the_references_on_that_page():
    # Sentences citing [1] fill the first pages and those citing [2] the last ones.
    slide = Slide("Mechanism", "content", [para(1) for _ in range(14)] + [para(2) for _ in range(14)], refs=REFS)
    prs = open_deck([slide])
    assert len(prs.slides) >= 2
    first, last = (named(s, "Footer references").text_frame.text for s in (prs.slides[0], prs.slides[-1]))
    assert first == "References: [1] Smith 2020"
    assert last.endswith("[2] Lee 2019") and "[1]" not in last
    link = named(prs.slides[0], "Footer references").text_frame.paragraphs[0].runs[1].hyperlink.address
    assert link == "https://example.org/smith"


def test_a_reference_without_a_link_is_still_listed_as_text():
    prs = open_deck([Slide("One", "content", [para(3)], refs=REFS)])
    runs = named(prs.slides[0], "Footer references").text_frame.paragraphs[0].runs
    assert "[3] Okafor 2021" in [r.text for r in runs]
    assert all(r.hyperlink.address is None for r in runs)


def test_a_slide_that_cites_nothing_has_no_references_line_but_a_page_number():
    prs = open_deck([Slide("Gap", "content", [Para("No verified finding in the papers read.")], refs=REFS)])
    names = [s.name for s in prs.slides[0].shapes]
    assert "Footer references" not in names
    assert named(prs.slides[0], "Page number").text_frame.text == "1 / 1"


def test_a_wide_figure_goes_above_the_text_and_a_tall_one_beside_it(tmp_path):
    wide = figure(png(tmp_path / "wide.png", 1600, 700))
    tall = figure(png(tmp_path / "tall.png", 700, 900), "Figure 3")
    slides = [
        Slide("Wide", "content", [para(1)], figure=wide, figure_ref=1, refs=REFS),
        Slide("Tall one", "content", [para(2)], figure=tall, figure_ref=2, refs=REFS),
        Slide("Tall two", "content", [para(1)], figure=tall, figure_ref=1, refs=REFS),
    ]
    prs = open_deck(slides)
    wide_pic, right_pic, left_pic = (named(s, "Figure") for s in prs.slides)
    text = lambda s: next(x for x in s.shapes if x.name == "Body")  # noqa: E731
    assert wide_pic.top >= text(prs.slides[0]).top + text(prs.slides[0]).height  # below the text
    assert right_pic.left > prs.slide_width // 2  # T2: text left, image right
    assert left_pic.left + left_pic.width < prs.slide_width // 2 + Emu(914400)  # T3: image left
    assert text(prs.slides[2]).left > left_pic.left + left_pic.width - Emu(914400 // 2)
    assert abs(wide_pic.width / wide_pic.height - 1600 / 700) < 0.01  # never stretched


def test_a_figure_is_credited_with_its_label_paper_and_licence(tmp_path):
    slide = Slide("MOA", "content", [para(1)], figure=figure(png(tmp_path / "f.png", 1600, 700)), figure_ref=1, refs=REFS)
    prs = open_deck([slide])
    credit = named(prs.slides[0], "Figure credit").text_frame
    assert credit.text == "Figure 2, [1] Smith 2020 - CC BY 4.0"
    runs = {r.text: r.hyperlink.address for r in credit.paragraphs[0].runs}
    assert runs["[1] Smith 2020"] == "https://example.org/smith"
    assert runs["CC BY 4.0"] == "https://creativecommons.org/licenses/by/4.0/"
    assert named(prs.slides[0], "Figure caption").text_frame.text == "Serum IL-23 fell after treatment."
    # the figure's paper counts as cited on the page, even when no sentence cites it
    only_figure = Slide("MOA", "content", [Para("No quote here.")], figure=figure(png(tmp_path / "g.png", 1600, 700)), figure_ref=2, refs=REFS)
    assert named(open_deck([only_figure]).slides[0], "Footer references").text_frame.text == "References: [2] Lee 2019"


def test_a_very_long_caption_is_cut_at_a_sentence_and_kept_whole_in_the_notes(tmp_path):
    caption = "First sentence about the figure. " + "More detail about panel B. " * 60
    slide = Slide("MOA", "content", [para(1)], figure=figure(png(tmp_path / "f.png", 1600, 700), caption=caption.strip()), figure_ref=1, refs=REFS)
    prs = open_deck([slide])
    shown = named(prs.slides[0], "Figure caption").text_frame.text
    assert shown.startswith("First sentence about the figure.") and shown.endswith("(caption continues in the speaker notes)")
    assert len(shown) < 520
    assert " ".join(caption.split()) in prs.slides[0].notes_slide.notes_text_frame.text


def test_an_unreadable_figure_leaves_the_slide_text_only(tmp_path):
    broken = tmp_path / "broken.png"
    broken.write_bytes(b"not an image")
    slides = [
        Slide("Broken", "content", [para(1)], figure=figure(str(broken)), figure_ref=1, refs=REFS),
        Slide("Missing", "content", [para(1)], figure=figure(str(tmp_path / "nope.png")), figure_ref=1, refs=REFS),
        Slide("Not downloaded", "content", [para(1)], figure=figure(None), figure_ref=1, refs=REFS),  # type: ignore[arg-type]
    ]
    prs = open_deck(slides)
    assert len(prs.slides) == 3
    assert all("Figure" not in [s.name for s in slide.shapes] for slide in prs.slides)


def test_a_figure_beats_a_visual_and_the_visual_still_appears_on_its_own_page(tmp_path):
    visual = Visual("doughnut_chart", [VisualItem("mouse", value=3)])
    slide = Slide("MOA", "content", [para(1)], visual=visual, figure=figure(png(tmp_path / "f.png", 1600, 700)), figure_ref=1, refs=REFS)
    prs = open_deck([slide])
    assert titles(prs) == ["MOA", "MOA (cont.)"]
    assert any(s.has_chart for s in prs.slides[1].shapes) and not any(s.has_chart for s in prs.slides[0].shapes)


def test_by_rule_tag_and_warning_are_on_the_slide_and_on_its_continuations():
    slide = Slide("Landscape", "content", [para(1) for _ in range(20)], by_rule=True, warning="Integrity check failed: see report.", refs=REFS)
    prs = open_deck([slide])
    assert len(prs.slides) > 1
    for s in prs.slides:
        body = texts(s)
        assert "Assembled by rule from the report - review before presenting" in body
        assert "Integrity check failed: see report." in body
    plain = open_deck([Slide("Plain", "content", [para(1)], refs=REFS)]).slides[0]
    assert not any("Assembled by rule" in t for t in texts(plain))


def test_kicker_is_under_the_title_on_the_first_page_only():
    slide = Slide("Mechanism", "content", [para(1) for _ in range(20)], kicker="Animal and cell evidence", refs=REFS)
    prs = open_deck([slide])
    assert "Animal and cell evidence" in texts(prs.slides[0])
    assert all("Animal and cell evidence" not in texts(s) for s in list(prs.slides)[1:])


def test_a_table_continues_with_its_header_row_repeated_and_loses_no_row():
    rows = [[f"Agent {i}", "IL-23", "psoriasis", f"[{1 + i % 3}]"] for i in range(40)]
    slide = Slide("Competitive landscape", "table", table=[["Drug/agent", "Target", "Indication", "Ref."], *rows], by_rule=True, refs=REFS)
    prs = open_deck([slide])
    assert len(prs.slides) > 1
    seen = []
    for number, s in enumerate(prs.slides, 1):
        grid = next(x for x in s.shapes if x.has_table).table
        assert [c.text for c in grid.rows[0].cells] == ["Drug/agent", "Target", "Indication", "Ref."]
        seen += [grid.cell(r, 0).text for r in range(1, len(grid.rows))]
        table_shape = next(x for x in s.shapes if x.has_table)
        assert table_shape.top + table_shape.height <= min(x.top for x in s.shapes if x.name.startswith("Footer"))
    assert seen == [r[0] for r in rows]
    assert titles(prs)[1] == "Competitive landscape (cont.)"


def test_a_table_footer_lists_the_references_cited_in_the_rows_on_that_page():
    slide = Slide("Landscape", "table", table=[["Drug", "Ref."], ["Agent A", "[2]"]], refs=REFS)
    footer = named(open_deck([slide]).slides[0], "Footer references").text_frame.text
    assert footer == "References: [2] Lee 2019"


def test_a_table_slide_with_no_rows_shows_its_statement():
    slide = Slide("Landscape", "table", [Para("No human trial named an agent - fill in.")], by_rule=True, refs=REFS)
    assert "No human trial named an agent - fill in." in " ".join(texts(open_deck([slide]).slides[0]))


def test_the_swot_matrix_has_four_cards_and_continues_when_items_do_not_fit():
    quads = []
    for name in ("Strengths", "Weaknesses", "Opportunities", "Threats"):
        quads += [Para(name, bold=True), *[para(1 + i % 3) for i in range(6)]]
    prs = open_deck([Slide("Target SWOT", "matrix", quads, by_rule=True, refs=REFS)])
    assert len(prs.slides) > 1
    first = " ".join(texts(prs.slides[0]))
    for name in ("Strengths", "Weaknesses", "Opportunities", "Threats"):
        assert name in first
    shown = [p.text for s in prs.slides for sh in s.shapes if sh.name == "Card text" for p in sh.text_frame.paragraphs
             if p.text.startswith("IL-23")]
    assert len(shown) == 24  # every item, once
    for s in prs.slides:
        assert all(sh.top + sh.height <= min(x.top for x in s.shapes if x.name.startswith("Footer")) for sh in s.shapes
                   if sh.name == "Card text")


def test_title_slide_has_the_title_and_subtitle_lines_and_no_page_number():
    prs = open_deck([Slide("IL-23 x TNF", "title", subtitle="Ayush Pawar | 30 Sep 2026\nWhat should I target?", warning="Check the report.")])
    body = texts(prs.slides[0])
    assert {"IL-23 x TNF", "Ayush Pawar | 30 Sep 2026", "What should I target?", "Check the report."} <= set(body)
    assert not any(s.name == "Page number" for s in prs.slides[0].shapes)


def test_document_properties_and_speaker_notes():
    slides = [Slide("IL-23 x TNF", "title", subtitle="Ayush Pawar | 30 Sep 2026\nQ?"),
              Slide("Summary", "content", [para(1)], notes=["E1: quote id", "E2"], refs=REFS)]
    prs = open_deck(slides)
    assert prs.core_properties.title == "IL-23 x TNF"
    assert prs.core_properties.author == "Ayush Pawar"
    assert prs.core_properties.comments == DECK_VERSION
    assert prs.slides[1].notes_slide.notes_text_frame.text == "E1: quote id\nE2"


def test_the_same_slides_give_the_same_deck(tmp_path):
    slides = every_layout() + [Slide("MOA", "content", [para(1)], figure=figure(png(tmp_path / "f.png", 1600, 700)), figure_ref=1, refs=REFS)]
    from lxml import etree

    def xml(deck):
        return [etree.tostring(s._element) for s in deck.slides]

    assert xml(open_deck(slides)) == xml(open_deck(slides))


def test_an_empty_slide_list_still_gives_a_valid_file():
    assert len(open_deck([]).slides) == 0


def test_importing_the_renderer_does_not_import_python_pptx():
    code = "import sys, bedrock.litreview.report.deck_render; assert 'pptx' not in sys.modules"
    assert subprocess.run([sys.executable, "-c", code], capture_output=True).returncode == 0
