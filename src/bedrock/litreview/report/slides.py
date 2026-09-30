"""A PowerPoint deck condensed from a finished Report. Code only - no model.

The deck is another renderer of the same structured report (stages/s8_render.py
renders JSON, Markdown and HTML), not a second piece of research: every
statement on a slide is a sentence the report kept under the citation guard,
a line the report's own code generated, or a paper's verbatim quote. Nothing is
summarised, reworded or added, so the deck cannot say more than the report.

What it leaves out is said on the title slide: the deck is condensed, and the
report is the record. Slides are split, never shrunk to fit: text that does not
fit continues on a numbered "(cont.)" slide, so nothing is cut off. No charts:
the report has no verified values laid out for one, and a chart drawn from
extracted numbers would state what no quote was checked for.

Needs python-pptx (`pip install -e ".[server]"`); imported only when a deck is
made, so the review itself never needs it.
"""

from __future__ import annotations

import io
import math
import re
from dataclasses import dataclass, field

from bedrock.litreview.models import Report, ReportSection
from bedrock.litreview.stages.s8_render import _numbers, _ref_map, _shown, lede

SLIDES_VERSION = "slides.v1"

CONDENSED = ("This deck is condensed from the full literature-review report. Every statement is a sentence the report "
             "kept, with the same citations; the verbatim quotes behind them are in the Evidence appendix. The report "
             "is the record: its evidence table, excluded papers and method are not repeated here.")
ABSENCE = ("No reported harm among these papers is not evidence of safety: the papers read may not have looked for, "
           "or been large or long enough to see, a given harm.")
PRECLINICAL = {"preclinical", "mechanism"}
_PRECLINICAL_NOTE = "Animal, cell and mechanism evidence - not results in people."

# Widescreen 13.33 x 7.5 in. Text is measured by a conservative width model
# (an average character half the font size wide: wider than Calibri, the
# deck's font, or Arial), so a slide is split before its text could overflow.
WIDTH_IN, HEIGHT_IN = 13.333, 7.5
BODY_WIDTH_IN = 12.1
_BOTTOM_IN = 0.55
_LINE = 1.2
_SPACE_AFTER = {0: 10, 1: 6}


def lines_needed(text: str, size: int, width_in: float = BODY_WIDTH_IN) -> int:
    per_line = max(1, int((width_in - 0.2) * 72 / (size * 0.5)))
    lines = 0
    for part in text.split("\n"):
        count, line = 1, 0
        for word in part.split(" "):
            count += (len(word) - 1) // per_line  # a word (a URL) wider than a line breaks across lines
            if line and line + 1 + len(word) > per_line:
                count, line = count + 1, len(word) % per_line
            else:
                line += (1 if line else 0) + len(word)
        lines += count
    return lines


def text_height_pt(text: str, size: int, level: int = 0, width_in: float = BODY_WIDTH_IN) -> float:
    return lines_needed(("    " * level) + text, size, width_in - 0.4 * level) * size * _LINE + _SPACE_AFTER[min(level, 1)]


def _top(kicker: bool, warning: bool) -> float:
    """Where a slide's body starts, in inches: under the title, the kicker line
    and the integrity warning when there is one."""
    return 1.3 + (0.45 if kicker else 0) + (0.6 if warning else 0)


def body_pt(kicker: bool = False, warning: bool = False) -> float:
    return (HEIGHT_IN - _top(kicker, warning) - _BOTTOM_IN) * 72 - 8


def plain(md: str) -> str:
    """Markdown the report's code wrote, as slide text: emphasis and link
    brackets dropped, whitespace collapsed. Words are never changed."""
    md = re.sub(r"<sub>|</sub>", "", md)
    md = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1", md)
    md = re.sub(r"<(https?://[^>\s]+)>", r"\1", md)
    md = re.sub(r"\*\*(.+?)\*\*", r"\1", md)
    md = re.sub(r"(?<![*\w])\*(?!\s)(.+?)(?<!\s)\*(?![*\w])", r"\1", md)
    return " ".join(md.split())


# --- what goes on the slides ---------------------------------------------------------------------------------


@dataclass
class Para:
    text: str
    level: int = 0
    bold: bool = False
    link: str | None = None
    muted: bool = False


@dataclass
class Slide:
    title: str
    paras: list[Para] = field(default_factory=list)
    table: list[list[str]] | None = None
    """Header row first."""
    kicker: str = ""
    """One line under the title: what kind of evidence the slide holds."""
    notes: list[str] = field(default_factory=list)
    size: int = 16
    warning: str = ""


def _section(report: Report, sid: str) -> ReportSection | None:
    return next((s for s in report.sections if s.id == sid), None)


def _bullets(md: str | None) -> list[str]:
    """The top-level "- " items of a generated section's Markdown."""
    return [plain(line[2:]) for line in (md or "").splitlines() if line.startswith("- ")]


def _first_paragraph(md: str | None) -> str:
    for block in (md or "").split("\n\n"):
        if block.strip() and not block.lstrip().startswith(("|", "- ", "#")):
            return plain(block)
    return ""


def _sentences(section: ReportSection, refs: dict[str, int], cited: list[str]) -> tuple[list[Para], list[str]]:
    """Kept sentences with their reference numbers, and a speaker-note line per
    sentence naming the evidence it rests on."""
    paras, notes = [], []
    kept = [s for s in section.sentences if s.kept]
    for i, s in enumerate(kept, 1):
        nums = _numbers(s, refs)
        marker = "".join(f"[{n}]" for n in nums)
        paras.append(Para(f"{s.text.rstrip()} {marker}".rstrip()))
        cited.extend(c for c in s.citations if c in refs and c not in cited)
        notes.append(f"Sentence {i}: " + ", ".join(f"{c} [{refs[c]}]" for c in s.citations if c in refs))
    return paras, notes


_SYNTH = re.compile(r"^(?P<label>.+?) — (?P<found>.+?)(?: Basis: (?P<basis>[^(]+?)\s*\((?P<why>.*?)\)\.)(?P<rest>.*)$")


def synthesis_rows(md: str | None) -> tuple[list[list[str]], list[str]]:
    """"How the studies line up" as table rows (measure, what was found, basis),
    plus the lines that are not a measure ("15 further measures ...")."""
    rows, other = [], []
    for item in _bullets(md):
        m = _SYNTH.match(item)
        if not m:
            other.append(item)
            continue
        found = m["found"].strip()
        rest = m["rest"].strip()
        rows.append([m["label"].strip(), f"{found} {rest}".strip(), f"{m['basis'].strip()} ({m['why'].strip()})"])
    return rows, other


# The limitations that most change how the Answer should be read, in order:
# whether the studies are the population asked about, which measures have no
# tested human effect, whether their numbers
# were printed wrongly, what was only read as an abstract, notices and
# preprints. The Limitations slide keeps all of them in the report's order.
_KEY_LIMITS = (r"^The question is about", r"^No tested effect in humans", r"interval printed|p-value printed|different estimates",
               r"abstract only", r"publication notice", r"preprint")


def key_limitations(bullets: list[str], limit: int = 3) -> list[str]:
    chosen = []
    for pattern in _KEY_LIMITS:
        chosen += [b for b in bullets if re.search(pattern, b) and b not in chosen]
    return chosen[:limit]


# --- pagination: split, never shrink ---------------------------------------------------------------------------


def _paginate(slide: Slide) -> list[Slide]:
    pages, current, used = [], [], 0.0
    for p in slide.paras:
        need = text_height_pt(p.text, slide.size, p.level)
        budget = body_pt(bool(slide.kicker), bool(slide.warning) and not pages)
        if current and used + need > budget:
            pages.append(current)
            current, used = [], 0.0
        current.append(p)
        used += need
    pages.append(current)
    out = []
    for i, paras in enumerate(pages):
        title = slide.title if i == 0 else f"{slide.title} (cont.)"
        out.append(Slide(title=title, paras=paras, kicker=slide.kicker, notes=slide.notes if i == 0 else
                         [f"Continues '{slide.title}'."], size=slide.size, warning=slide.warning if i == 0 else ""))
    return out


def row_height_pt(row: list[str], widths: list[float], size: int = 12) -> float:
    return max(lines_needed(cell, size, w) for cell, w in zip(row, widths)) * size * _LINE + 10


def _paginate_table(slide: Slide) -> list[Slide]:
    head, rows = slide.table[0], slide.table[1:]
    widths = _COLUMN_WIDTHS.get(len(head), [BODY_WIDTH_IN / len(head)] * len(head))
    tail = sum(text_height_pt(p.text, 12) for p in slide.paras)
    pages, current, used = [], [], row_height_pt(head, widths)
    for i, row in enumerate(rows):
        need = row_height_pt(row, widths) + (tail if i == len(rows) - 1 else 0)
        if current and used + need > body_pt(bool(slide.kicker)):
            pages.append(current)
            current, used = [], row_height_pt(head, widths)
        current.append(row)
        used += need
    pages.append(current)
    return [Slide(title=slide.title if i == 0 else f"{slide.title} (cont.)", table=[head, *rows_], kicker=slide.kicker,
                  paras=slide.paras if i == len(pages) - 1 else [], notes=slide.notes, size=12)
            for i, rows_ in enumerate(pages)]


_COLUMN_WIDTHS = {3: [3.0, 5.6, 3.5]}


# --- the deck --------------------------------------------------------------------------------------------------------


def outline(report: Report, source: str = "") -> list[Slide]:
    """The slides, as data: what the deck says, independent of python-pptx.
    Deterministic - the same report gives the same slides."""
    refs, quotes = _ref_map(report)
    cited: list[str] = []
    slides: list[Slide] = []
    stats = report.stats
    interp = report.interpretation
    warning = ("This report failed its integrity check: a quote or field did not match its paper. Do not rely on it "
               "until it is fixed.") if report.integrity else ""

    title_paras = [Para(CONDENSED)]
    if source:
        title_paras.append(Para(f"Source report: {source}", muted=True))
    slides.append(Slide(title=report.question, paras=title_paras, kicker="Literature review - condensed deck",
                        notes=[f"Deck {SLIDES_VERSION}, made by code from the stored report; no model was used.",
                               "Models behind the report: " + ", ".join(f"{k} {v}" for k, v in
                                                                         sorted(report.model_versions.items()))],
                        warning=warning, size=18))

    scope = [Para(f"Question: {report.question}", bold=True)]
    read_as = [f"{str(interp.get('question_type', '')).replace('_', ' ')}", f"goal: {interp.get('goal', '')}"]
    if interp.get("subject"):
        read_as.append(f"subject: {interp['subject']}")
    if interp.get("population"):
        read_as.append(f"population: {interp['population']}")
    scope.append(Para("Read as: " + "; ".join(x for x in read_as if x.strip())))
    full = stats.get("papers_read", 0) - stats.get("papers_abstract_only", 0)
    gave = len(report.references)
    scope.append(Para(f"Papers: {stats.get('papers_collected', 0)} collected; {stats.get('gate_pass', 0)} concern the "
                      f"subject exactly; {stats.get('papers_read', 0)} read ({full} in full text, "
                      f"{stats.get('papers_abstract_only', 0)} from the abstract only); {gave} of them gave verified "
                      "findings, and the counts below are of those."))
    glance = _section(report, "glance")
    if glance is not None and _first_paragraph(glance.body_markdown):
        scope.append(Para(_first_paragraph(glance.body_markdown)))
    slides.append(Slide(title="Question and scope", paras=scope, notes=[lede(report)]))

    answer = _section(report, "answer")
    if answer is not None:
        paras, notes = _sentences(answer, refs, cited)
        if not paras:
            paras = [Para(plain(answer.note or "No sentence of the answer survived the citation check."))]
        elif answer.note:
            paras.append(Para(plain(answer.note), muted=True))
        gaps = key_limitations(_bullets((_section(report, "gaps") or
                                        ReportSection(id="", title="", kind="generated")).body_markdown))
        if gaps:
            paras.append(Para("Key limitations", bold=True))
            paras += [Para(g, level=1) for g in gaps]
        slides += _paginate(Slide(title="Answer", paras=paras, notes=notes, warning=warning))

    synth = _section(report, "synthesis")
    if synth is not None and synth.body_markdown:
        rows, other = synthesis_rows(synth.body_markdown)
        if rows:
            table = Slide(title="How the studies line up", table=[["Measure", "What the studies found", "Basis"], *rows],
                          kicker="Grouped by what was measured; publications from one trial count once",
                          paras=[Para(o, muted=True) for o in other],
                          notes=["Basis says what kind of evidence stands behind each line - not a quality or "
                                 "certainty rating. Findings the requested treatment was not responsible for, or "
                                 "that compare subgroups or arms that all received it, are listed as not counted."])
            slides += _paginate_table(table)

    for section in report.sections:
        # Safety always has a slide: "nothing found" is itself the finding a
        # reader of a deck needs to see, with why it is not reassurance.
        if section.kind != "written" or section.id == "answer" or (not _shown(section) and section.id != "safety"):
            continue
        paras, notes = _sentences(section, refs, cited)
        if section.id == "safety":
            if not paras:
                paras = [Para("No verified safety finding in the papers read.")]
            paras.append(Para(ABSENCE, muted=True))
        if not paras:
            continue
        if section.note and _shown(section):
            paras.append(Para(plain(section.note), muted=True))
        kicker = _PRECLINICAL_NOTE if section.id in PRECLINICAL else ""
        slides += _paginate(Slide(title=section.title, paras=paras, notes=notes, kicker=kicker))

    gaps = _section(report, "gaps")
    if gaps is not None:
        items = _bullets(gaps.body_markdown)
        empty = [s.title for s in report.sections if s.kind == "written" and not _shown(s) and s.id != "safety"]
        if empty:
            items.append("No verified finding for: " + ", ".join(empty) + ".")
        if items:
            slides += _paginate(Slide(title="Limitations and open questions", paras=[Para(i) for i in items],
                                      size=14))

    ref_paras = [Para(f"[{r.number}] {plain(r.citation)}", link=r.url) for r in report.references]
    if ref_paras:
        slides += _paginate(Slide(title="References", paras=ref_paras, size=12,
                                  notes=["Each reference links to its record (Europe PMC or the publisher)."]))

    quote_paras = [Para(f"{c} [{refs[c]}] “{' '.join(quotes[c].split())}”") for c in cited]
    if quote_paras:
        slides += _paginate(Slide(title="Appendix: evidence quotes", paras=quote_paras, size=12,
                                  kicker="Each paper's own words, exactly as verified; cited above by these ids"))
    return slides


def to_pptx(report: Report, source: str = "") -> bytes:
    from pptx import Presentation  # ImportError when python-pptx is missing: the caller says how to install it
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    ink, muted, accent, warn = (RGBColor(0x1F, 0x29, 0x37), RGBColor(0x5B, 0x65, 0x73), RGBColor(0x1D, 0x4E, 0x89),
                                RGBColor(0xA1, 0x26, 0x1B))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH_IN), Inches(HEIGHT_IN)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = report.question[:250]
    prs.core_properties.subject = "Condensed literature-review report"
    prs.core_properties.author = "Bedrock literature review"
    prs.core_properties.last_modified_by = "Bedrock literature review"
    prs.core_properties.comments = SLIDES_VERSION

    def text_box(slide, left, top, width, height):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        return frame

    def run(paragraph, text, size, color, bold=False, link=None):
        r = paragraph.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
        if link:
            r.hyperlink.address = link
        return r

    outline_ = outline(report, source)
    for number, spec in enumerate(outline_, 1):
        slide = prs.slides.add_slide(blank)
        top = 0.35
        title_size = 30 if number == 1 else 26
        frame = text_box(slide, 0.6, top, 12.1, 1.0 if number > 1 else 1.6)
        run(frame.paragraphs[0], spec.title, title_size if len(spec.title) < 90 else 22, accent, bold=True)
        top = 1.3 if number > 1 else 2.05
        if spec.kicker:
            run(text_box(slide, 0.6, top, 12.1, 0.4).paragraphs[0], spec.kicker, 13, muted)
            top += 0.45
        if spec.warning:
            run(text_box(slide, 0.6, top, 12.1, 0.6).paragraphs[0], spec.warning, 13, warn, bold=True)
            top += 0.6
        if spec.table:
            widths = _COLUMN_WIDTHS.get(len(spec.table[0]), [12.1 / len(spec.table[0])] * len(spec.table[0]))
            shape = slide.shapes.add_table(len(spec.table), len(spec.table[0]), Inches(0.6), Inches(top),
                                           Inches(12.1), Inches(0.4 * len(spec.table)))
            for c, w in enumerate(widths):
                shape.table.columns[c].width = Inches(w)
            for r_, row in enumerate(spec.table):
                for c, value in enumerate(row):
                    cell = shape.table.cell(r_, c)
                    cell.text_frame.word_wrap = True
                    p = cell.text_frame.paragraphs[0]
                    run(p, value, 12, ink, bold=r_ == 0)
            top += sum(row_height_pt(row, widths) for row in spec.table) / 72 + 0.15
        if spec.paras:
            frame = text_box(slide, 0.6, top, 12.1, HEIGHT_IN - top - 0.5)
            for i, p in enumerate(spec.paras):
                para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                para.level = p.level
                para.space_after = Pt(_SPACE_AFTER[min(p.level, 1)])
                prefix = "• " if p.level else ""
                run(para, prefix + p.text, spec.size - (1 if p.muted else 0), muted if p.muted else ink, bold=p.bold,
                    link=p.link)
        foot = text_box(slide, 0.6, HEIGHT_IN - 0.45, 12.1, 0.3).paragraphs[0]
        run(foot, f"Condensed from the full report · {number} / {len(outline_)}", 10, muted)
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(spec.notes)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
