"""Draws the target-proposal deck (`deck_model.Slide` objects) as a PowerPoint
file. This is a python-pptx port of ArcDeck's slide format - the same canvas,
margins, title bar, image layouts, visual elements and font sizes - in
Propeller's template colours and fonts. It draws what it is given and decides
nothing about content: what a slide says is `report/slides.py`'s job, and every
word on it was made there from the stored report. Nothing here asks a model
anything or touches the network.

Three rules shape the code, because a deck that looks wrong in front of a
client is worse than one that is a page longer:

* **Split, never shrink.** Text keeps the size the slide asked for. A slide
  whose text does not fit continues on a "(cont.)" slide; a table continues on
  one with its header row repeated. Heights come from `deck_model.lines_needed`
  - a conservative width model, so a page is split before its text could reach
  the footer - using the real width of the text box each layout gives the text.
* **Only what is on the page is in its footer.** Each page's footer lists the
  references whose numbers appear on that page, each linked to its paper, so a
  citation can be followed from a printed or projected page.
* **A figure is credited where it is shown.** Its label, the paper it came from
  and its licence are printed under it, always.

Needs python-pptx (`pip install -e ".[server]"`); it is imported inside
`render()`, so the review itself never needs it. Pillow, which python-pptx
depends on, reads each figure's size to choose its layout.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from bedrock.litreview.report.deck_model import (
    _SPACE_AFTER,
    DECK_VERSION,
    HEIGHT_IN,
    WIDTH_IN,
    FigureAsset,
    Para,
    RefMark,
    Slide,
    Visual,
    lines_needed,
)

# --- Propeller's theme on ArcDeck's "ocean" structure -------------------------------------------------------

FH, FB = "Franklin Gothic Demi", "Corbel"
NAVY, DEEP, PRI, TEAL, ACC, ACC_LT = "335B74", "26475A", "2683C6", "1CADE4", "27CED7", "A6E7EF"
WH, BG, CARD, TX, TX_M, TX_L = "FFFFFF", "F4F6F7", "FFFFFF", "1E293B", "475569", "94A3B8"
RED, GRN, DIV, ZEBRA = "DC2626", "42BA97", "DFE3E5", "E8EEF2"
VE_COLS = [PRI, NAVY, TEAL, GRN]
CHART_COLS = [PRI, TEAL, ACC, NAVY, GRN, "62A39F"]
OUTCOME_COLS = {"improved": GRN, "no_effect": TX_L, "worsened": RED, "mixed": PRI, "association": NAVY}
"""What a bar means never changes with its position: an improved finding is
always green, a worsened one always red."""
TONE_COLS = {"positive": GRN, "negative": TX_L, "neutral": PRI}
SWOT_COLS = [GRN, RED, PRI, NAVY]

# --- geometry, in inches (ArcDeck: SW, SH, ML, CW) ----------------------------------------------------------

SW, SH = WIDTH_IN, HEIGHT_IN
ML = 0.6
CW = SW - 2 * ML
"""Width of the text column. Text boxes sit 0.1 in left of ML and carry a 0.1 in
inset, so the text itself starts exactly on the margin the title uses."""
_TITLE_SIZE = 24
_FOOT_BOTTOM, _FOOT_GAP, _FOOT_MIN = 0.15, 0.12, 0.24
_FOOT_W = CW - 1.1
_IMG_COL = 5.2
_IMG_GAP = 0.35
_FIG_MIN_H = 2.2
_WIDE = 1.6
"""A figure at least this much wider than tall goes above the text (ArcDeck T4);
a squarer one goes beside it (T2 on the right, T3 on the left, alternating)."""
_TABLE_SIZE = 12
_CAPTION_LIMIT = 450
_BY_RULE = "Assembled by rule from the report - review before presenting"
_CITE = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")


def _lines(text: str, size: int, width_in: float) -> int:
    return lines_needed(text, size, width_in)


def _need(text: str, size: int, width: float) -> float:
    """Height in inches of `text` in a box with no inset that is `width` wide."""
    return _lines(text, size, width + 0.2) * size * 1.2 / 72


def _numbers_in(text: str) -> set[int]:
    return {int(n) for m in _CITE.finditer(text) for n in m.group(1).split(",")}


def _para_refs(p: Para) -> set[int]:
    return set(p.refs) | _numbers_in(p.text)


# --- what each page needs, worked out before anything is drawn -----------------------------------------------


@dataclass
class _Page:
    """One drawn slide: a source slide, or one page of a split one."""

    src: Slide
    kind: str  # "title" | "flow" | "matrix"
    title: str
    first: bool = True
    paras: list[Para] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    refs: list[int] = field(default_factory=list)
    top: float = 0.0
    bottom: float = 0.0
    text_w: float = CW
    text_h: float = 0.0
    visual: Visual | None = None
    figure: FigureAsset | None = None
    fig_px: tuple[int, int] = (1, 1)
    variant: str = ""
    groups: list[tuple[str, list[Para]]] = field(default_factory=list)
    extra_notes: list[str] = field(default_factory=list)


@dataclass
class _Deck:
    side_figures: int = 0
    """Side-image slides so far: even ones put the image right, odd ones left."""


def _header_heights(title: str, kicker: str, warning: str, by_rule: bool) -> list[float]:
    return [
        max(0.5, _need(title, _TITLE_SIZE, CW) + 0.1),
        _need(kicker, 12, CW) + 0.04 if kicker else 0.0,
        _need(warning, 12, CW) + 0.04 if warning else 0.0,
        _need(_BY_RULE, 10, CW) + 0.04 if by_rule else 0.0,
    ]


def _body_top(title: str, kicker: str, warning: str, by_rule: bool) -> float:
    return 0.25 + sum(_header_heights(title, kicker, warning, by_rule)) + 0.12


def _footer_marks(numbers: set[int], marks: dict[int, RefMark]) -> list[tuple[int, RefMark | None]]:
    return [(n, marks.get(n)) for n in sorted(numbers)]


def _footer_text(items: list[tuple[int, RefMark | None]]) -> str:
    return "References: " + " · ".join(f"[{n}] {m.short}" if m else f"[{n}]" for n, m in items)


def _foot_h(items: list[tuple[int, RefMark | None]]) -> float:
    return max(_FOOT_MIN, _need(_footer_text(items), 10, _FOOT_W) + 0.06) if items else _FOOT_MIN


def _bottom(items: list[tuple[int, RefMark | None]]) -> float:
    """Where a page's body must end: just above its footer, which is as tall as
    the references on the page need."""
    return SH - _FOOT_BOTTOM - _foot_h(items) - _FOOT_GAP


def _bulleted(p: Para) -> bool:
    return p.level > 0 or not p.bold  # a bold top-level paragraph is a sub-heading


def _para_size(p: Para, base: int) -> int:
    size = base - 2 if p.level else base
    return max(8, size - 1 if p.muted else size)


def _para_h(p: Para, base: int, text_w: float, bullets: bool) -> float:
    """Inches a paragraph needs in a text box whose text column is `text_w` wide."""
    indent = (0.25 if bullets and _bulleted(p) else 0.0) + 0.4 * p.level
    size = _para_size(p, base)
    return (_lines(p.text, size, text_w + 0.2 - indent) * size * 1.2 + _SPACE_AFTER[min(p.level, 1)]) / 72


def _split_para(p: Para, base: int, text_w: float, bullets: bool, room: float) -> tuple[Para, Para | None]:
    """A single paragraph taller than a whole page is split between words - the
    only way text is ever cut, and nothing is lost: the rest continues."""
    words = p.text.split(" ")
    low, high = 1, len(words)
    while low < high:  # the most words that fit
        mid = (low + high + 1) // 2
        if _para_h(Para(" ".join(words[:mid]), p.level, p.bold, None, p.muted), base, text_w, bullets) <= room:
            low = mid
        else:
            high = mid - 1
    if low >= len(words):
        return p, None
    head = Para(" ".join(words[:low]), p.level, p.bold, p.link, p.muted, list(p.refs))
    tail = Para(" ".join(words[low:]), p.level, p.bold, p.link, p.muted, list(p.refs))
    return head, tail


def _column_widths(table: list[list[str]]) -> list[float]:
    """Column widths that add up to the slide's text width: wider for columns
    whose cells hold more text, and never narrower than 0.8 in."""
    weights = []
    for c in range(len(table[0])):
        cells = [row[c] if c < len(row) else "" for row in table]
        body = cells[1:] or [""]
        longest = max((len(w) for cell in cells for w in cell.split()), default=1)
        weights.append(min(max(sum(map(len, body)) / len(body), longest * 1.1, 6), 34))
    widths = [CW * w / sum(weights) for w in weights]
    short = [i for i, w in enumerate(widths) if w < 0.8]
    if short:
        rest = [i for i in range(len(widths)) if i not in short]
        spare = CW - 0.8 * len(short)
        total = sum(widths[i] for i in rest) or 1
        widths = [0.8 if i in short else widths[i] * spare / total for i in range(len(widths))]
    return widths


def _row_h(row: list[str], widths: list[float], size: int = _TABLE_SIZE) -> float:
    lines = max((_lines(cell, size, w + 0.04) for cell, w in zip(row, widths)), default=1)
    return lines * size * 1.2 / 72 + 0.1


def _figure_px(figure: FigureAsset | None) -> tuple[int, int] | None:
    """The downloaded image's pixel size, or None when there is no readable image:
    a figure that cannot be drawn is left out and the slide still renders."""
    if figure is None or not figure.path:
        return None
    try:
        from PIL import Image

        with Image.open(figure.path) as image:
            width, height = image.size
    except Exception:  # noqa: BLE001 - a missing, corrupt or unsupported file never fails a deck
        return None
    return (width, height) if width > 0 and height > 0 else None


def _caption(figure: FigureAsset) -> tuple[str, bool]:
    """The paper's own caption, whole when short. A very long one is cut at a
    sentence end and says so; the speaker notes keep all of it."""
    text = " ".join(figure.caption.split())
    if len(text) <= _CAPTION_LIMIT:
        return text, False
    cut = text[:_CAPTION_LIMIT]
    end = max(cut.rfind(". "), cut.rfind("? "), cut.rfind("! "))
    cut = cut[: end + 1] if end >= 80 else cut[: cut.rfind(" ")] + " ..." if " " in cut else cut
    return cut + " (caption continues in the speaker notes)", True


def _visual_bounds(v: Visual, alone: bool = False) -> tuple[float, float]:
    """The shortest and tallest a visual may be, caption line included.
    Cards and stacked rows are exactly as tall as their text needs; charts take
    a floor of 2.0 in (ArcDeck's cap) and grow into spare space, further when
    the slide has no text (`alone`)."""
    cap = _need(v.caption, 11, CW) + 0.1 if v.caption else 0.0
    items = v.items[:8]
    n = max(1, len(items))
    if v.type == "stats_callout":
        return 1.8 + cap, 1.8 + cap
    if v.type == "callout_cards":
        n = min(n, 6)
        w = (CW - (n - 1) * 0.25) / n - 0.35
        need = max((_need(i.label, 12, w) + 0.08 + _need(i.sublabel, 11, w) + 0.3 for i in items[:6]), default=1.1)
        return min(max(need, 1.1), 2.8) + cap, min(max(need, 1.1), 2.8) + cap
    if v.type == "comparison":
        n = min(n, 5)
        w = (CW - (n - 1) * 0.2) / n - 0.2
        head = max((_need(i.label, 13, w) + 0.14 for i in items[:5]), default=0.45)
        need = max(
            (
                max(0.45, head) + 0.1 + (0.55 if i.value is not None else 0)
                + sum(_need(line, 11, w) for line in (i.sublabel or " ").split("\n")) + 0.15
                for i in items[:5]
            ),
            default=1.6,
        )
        return min(max(need, 1.6), 3.4) + cap, min(max(need, 1.6), 3.4) + cap
    if v.type == "stacked_list":
        rows = _stack_rows(items)
        return sum(rows) + 0.07 * (len(rows) - 1) + cap, sum(rows) + 0.07 * (len(rows) - 1) + cap
    return 2.0 + cap, (4.4 if alone else 3.2) + cap  # bar_chart, doughnut_chart


_STACK_LABEL_W, _STACK_SUB_W = 3.3, CW - 4.4


def _stack_rows(items) -> list[float]:
    return [
        max(0.38, max(_need(i.label, 13, _STACK_LABEL_W), _need(i.sublabel, 12, _STACK_SUB_W)) + 0.14)
        for i in items[:8]
    ]


def _plan_flow(s: Slide, deck: _Deck) -> list[_Page]:
    """Split a content, table or references slide into pages and decide, for each,
    which paragraphs and table rows it holds, how tall they are and how tall
    its footer must be. The footer depends on which references are on the page
    and the page depends on the footer, so each page is filled greedily and the
    footer re-measured after every addition."""
    marks = {m.number: m for m in s.refs}
    bullets = s.layout != "references"
    footer_refs = s.layout != "references"
    paras = list(s.paras)
    header = s.table[0] if s.table and s.layout == "table" else None
    rows = [list(r) for r in s.table[1:]] if header else []
    widths = _column_widths(s.table) if header else []
    figure = visual = None
    variant, px = "", (1, 1)
    if s.layout == "content":
        px_ = _figure_px(s.figure)
        if px_:
            figure, px = s.figure, px_
            variant = "top" if px[0] / px[1] >= _WIDE else ("right" if deck.side_figures % 2 == 0 else "left")
            deck.side_figures += variant != "top"
        elif s.visual and s.visual.items:
            visual = s.visual
    pages: list[_Page] = []
    while True:
        first = not pages
        title = s.title if first else f"{s.title} (cont.)"
        top = _body_top(title, s.kicker if first else "", s.warning, s.by_rule)
        text_w = CW - (_IMG_COL + _IMG_GAP if figure and variant != "top" and first else 0)
        numbers: set[int] = set()
        reserved = 0.0
        credit_extra: list[str] = []
        if first and figure:
            if s.figure_ref is not None:
                numbers.add(s.figure_ref)
            cap, cut = _caption(figure)
            col = CW if variant == "top" else _IMG_COL
            below = _need(cap, 9, col) + 0.04 + _need(_credit_text(figure, s.figure_ref, marks), 8, col) + 0.08
            if cut:
                credit_extra.append(f"{figure.label} caption in full: {' '.join(figure.caption.split())}")
            reserved = (min(_FIG_MIN_H, CW * px[1] / px[0]) + below + 0.2) if variant == "top" else 0.0
        elif first and visual:
            reserved = _visual_bounds(visual)[0] + 0.2

        def room(nums: set[int], used: float) -> float:
            return _bottom(_footer_marks(nums, marks) if footer_refs else []) - top - reserved - used

        placed: list[Para] = []
        used = 0.0
        while paras:
            p = paras[0]
            h = _para_h(p, s.size, text_w, bullets)
            nums = numbers | _para_refs(p)
            keep, keep_nums = h, nums  # what must fit for p to stay: a heading is kept with the paragraph after it
            if p.bold and p.level == 0 and len(paras) > 1:
                keep, keep_nums = h + _para_h(paras[1], s.size, text_w, bullets), nums | _para_refs(paras[1])
            if keep <= room(keep_nums, used) or (not placed and h <= room(nums, used)):
                placed.append(paras.pop(0))
                numbers, used = nums, used + h
            elif not placed:  # an empty page takes what comes next, split if need be, so the loop always advances
                head, tail = _split_para(p, s.size, text_w, bullets, max(room(nums, used), 0.3))
                paras[0:1] = [tail] if tail else []
                placed.append(head)
                numbers |= _para_refs(head)
                used += _para_h(head, s.size, text_w, bullets)
                break
            else:
                break
        page_rows: list[list[str]] = []
        if header and not paras:
            gap = 0.12 if placed else 0.0
            table_h = _row_h(header, widths) + gap
            while rows:
                h = _row_h(rows[0], widths)
                nums = numbers | _numbers_in(" ".join(rows[0]))
                if table_h + h <= room(nums, used) or not page_rows and not placed:
                    page_rows.append(rows.pop(0))
                    numbers, table_h = nums, table_h + h
                else:
                    break
            if page_rows:
                used += table_h
        shown = _footer_marks(numbers, marks) if footer_refs else []
        pages.append(
            _Page(
                s, "flow", title, first, placed, page_rows, [n for n, _ in shown], top, _bottom(shown), text_w,
                sum(_para_h(p, s.size, text_w, bullets) for p in placed),
                visual if first else None, figure if first else None, px, variant if first else "",
                extra_notes=credit_extra if first else [],
            )
        )
        if not paras and not rows:
            break
    if figure and s.visual and s.visual.items:  # both given: the figure takes page 1, the visual its own page
        top = _body_top(f"{s.title} (cont.)", "", s.warning, s.by_rule)
        pages.append(_Page(s, "flow", f"{s.title} (cont.)", False, [], [], [], top, _bottom([]), CW, 0.0, s.visual))
    return pages


def _card_h(items: list[Para], size: int, text_w: float, show_empty: bool) -> float:
    if not items and not show_empty:
        return 0.0  # a card with nothing left to say is not repeated on a continuation slide
    return 0.5 + sum(_para_h(p, size, text_w, True) for p in items) + (0.1 if items else 0.25)


def _matrix_rows(cut: list[tuple[str, list[Para]]], size: int, text_w: float, first: bool) -> list[float]:
    """Height of each row of two cards: as tall as the taller card."""
    return [max(_card_h(items, size, text_w, first) for _, items in cut[r : r + 2]) for r in range(0, len(cut), 2)]


def _matrix_height(rows: list[float]) -> float:
    shown = [h for h in rows if h]
    return sum(shown) + 0.15 * max(len(shown) - 1, 0)


def _plan_matrix(s: Slide) -> list[_Page]:
    """The 2x2 cards (Strengths, Weaknesses, Opportunities, Threats). Each bold
    top-level paragraph names a card and the paragraphs after it are its items.
    Cards hold the same number of items on a page; items that do not fit
    continue on a "(cont.)" slide."""
    marks = {m.number: m for m in s.refs}
    groups: list[tuple[str, list[Para]]] = []
    for p in s.paras:
        if p.bold and p.level == 0:
            groups.append((p.text, []))
        else:
            if not groups:
                groups.append(("", []))
            groups[-1][1].append(p)
    if not groups and s.visual:  # cards given as visual items: label = card name, sublabel lines = items
        groups = [(i.label, [Para(line) for line in i.sublabel.split("\n") if line]) for i in s.visual.items]
    size = min(s.size, 12)
    card_w = (CW - 0.25) / 2
    text_w = card_w - 0.3
    pages: list[_Page] = []
    offset = 0
    longest = max((len(g[1]) for g in groups), default=0)
    while True:
        first = not pages
        title = s.title if first else f"{s.title} (cont.)"
        top = _body_top(title, s.kicker if first else "", s.warning, s.by_rule)

        def layout(k: int):
            cut = [(label, items[offset : offset + k]) for label, items in groups]
            shown = _footer_marks({n for _, items in cut for p in items for n in _para_refs(p)}, marks)
            return cut, shown, _matrix_rows(cut, size, text_w, first)

        best = 1
        for k in range(1, max(longest - offset, 1) + 1):
            _, shown, rows = layout(k)
            if _matrix_height(rows) <= _bottom(shown) - top:
                best = k
            else:
                break
        cut, shown, _ = layout(best)
        pages.append(_Page(s, "matrix", title, first, refs=[n for n, _ in shown], top=top, bottom=_bottom(shown), groups=cut))
        offset += best
        if offset >= longest:
            return pages


def _plan_title(s: Slide, deck: _Deck) -> list[_Page]:
    """The title slide has no footer. Any paragraphs given to it are drawn small
    under the subtitle while they fit and continue on a content slide after."""
    page = _Page(s, "title", s.title)
    room = SH - 0.5 - (0.8 if s.warning else 0) - _title_block_bottom(s)
    placed, used = [], 0.0
    paras = list(s.paras)
    while paras:
        h = _para_h(paras[0], 14, 11.3, True)
        if used + h > room and placed:
            break
        placed.append(paras.pop(0))
        used += h
    page.paras, page.text_h = placed, used
    pages = [page]
    if paras:
        rest = Slide(s.title, "content", paras, refs=s.refs, notes=s.notes, warning=s.warning, size=s.size)
        pages += _plan_flow(rest, deck)
    return pages


def _title_block_bottom(s: Slide) -> float:
    title_h = _need(s.title, 36, 11.3) + 0.1
    sub = s.subtitle.split("\n") if s.subtitle else []
    sub_h = sum(_need(t, 16 if i == 0 else 14, 11.3) + 0.08 for i, t in enumerate(sub))
    return 1.3 + title_h + 0.15 + 0.03 + 0.3 + sub_h + 0.2


def _credit_text(f: FigureAsset, ref: int | None, marks: dict[int, RefMark]) -> str:
    source = ""
    if ref is not None:
        source = f"[{ref}] {marks[ref].short}" if ref in marks else f"[{ref}]"
    return ", ".join(x for x in (f.label, source) if x) + (f" - {f.licence}" if f.licence else "")


# --- drawing -----------------------------------------------------------------------------------------------


def render(slides: list[Slide]) -> bytes:
    """The deck as .pptx file bytes. Each `Slide` becomes one or more PowerPoint
    slides; the same slides give the same deck."""
    from pptx import Presentation
    from pptx.util import Inches

    deck, pages = _Deck(), []
    for s in slides:
        if s.layout == "title":
            pages += _plan_title(s, deck)
        elif s.layout == "matrix":
            pages += _plan_matrix(s)
        else:
            pages += _plan_flow(s, deck)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    first = slides[0] if slides else None
    props = prs.core_properties
    props.title = (first.title if first else "Target proposal")[:250]
    props.author = props.last_modified_by = (
        first.subtitle.split("\n")[0].split(" | ")[0].strip() if first and first.subtitle else ""
    ) or "Bedrock literature review"
    props.subject = "Target proposal"
    props.comments = DECK_VERSION
    for number, page in enumerate(pages, 1):
        _draw_page(prs, page, number, len(pages))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _rgb(hex_: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_)


def _on(fill: str) -> str:
    """White or dark text, whichever reads better on this fill."""

    def lum(h: str) -> float:
        parts = [int(h[i : i + 2], 16) / 255 for i in (0, 2, 4)]
        lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in parts]
        return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]

    ratio = lambda a, b: (max(lum(a), lum(b)) + 0.05) / (min(lum(a), lum(b)) + 0.05)  # noqa: E731
    return WH if ratio(fill, WH) >= ratio(fill, TX) else TX


def _rect(slide, x, y, w, h, fill, *, rounded=False, line=None, shadow=False):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt
    from lxml import etree

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if rounded:
        shape.adjustments[0] = min(0.5, 0.05 / max(min(w, h), 0.01))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(0.5)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False  # no theme shadow; the soft one below is ArcDeck's
    if shadow:
        effects = shape._element.spPr.find(qn("a:effectLst"))
        outer = etree.SubElement(effects, qn("a:outerShdw"), blurRad="76200", dist="25400", dir="2700000", algn="tl", rotWithShape="0")
        etree.SubElement(etree.SubElement(outer, qn("a:srgbClr"), val="000000"), qn("a:alpha"), val="12000")
    return shape


def _frame(slide, x, y, w, h, *, inset=0.0, anchor="top"):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(inset)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    return tf


def _run(par, text, size, color, *, bold=False, italic=False, font=FB, link=None):
    from pptx.util import Pt

    r = par.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
    r.font.color.rgb = _rgb(color)
    if link:
        r.hyperlink.address = link
    return r


def _label(slide, x, y, w, h, text, size, color, *, bold=False, italic=False, font=FB, align="left", anchor="top"):
    """One piece of text in its own box (no inset), e.g. a card label."""
    from pptx.enum.text import PP_ALIGN

    tf = _frame(slide, x, y, w, h, anchor=anchor)
    for i, line in enumerate(text.split("\n")):  # a line break starts a new paragraph, so alignment holds on every line
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        _run(par, line, size, color, bold=bold, italic=italic, font=font)
    return tf


def _bullet(par, level, on, color):
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches((0.25 if on else 0.0) + 0.4 * level))))
    pPr.set("indent", str(-int(Inches(0.25))) if on else "0")
    if on:
        etree.SubElement(etree.SubElement(pPr, qn("a:buClr")), qn("a:srgbClr"), val=color)
        etree.SubElement(pPr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(pPr, qn("a:buChar"), char="•" if level == 0 else "–")
    else:
        etree.SubElement(pPr, qn("a:buNone"))


def _write_paras(tf, paras, base, *, bullets=True, color=TX, muted=TX_M, bullet=PRI):
    from pptx.util import Pt

    for i, p in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(_SPACE_AFTER[min(p.level, 1)])
        _bullet(par, p.level, bullets and _bulleted(p), bullet)
        _run(par, p.text, _para_size(p, base), muted if p.muted else color, bold=p.bold, link=p.link)


def _draw_page(prs, page: _Page, number: int, total: int) -> None:
    from pptx.util import Inches

    s = page.src
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(NAVY if page.kind == "title" else BG)
    if page.kind == "title":
        _draw_title_page(slide, page)
    else:
        _rect(slide, 0, 0, SW, 0.05, PRI)
        _draw_header(slide, page)
        if page.kind == "matrix":
            _draw_matrix(slide, page)
        else:
            _draw_flow(slide, page)
        _draw_footer(slide, page, number, total)
    if s.notes or page.extra_notes:
        slide.notes_slide.notes_text_frame.text = "\n".join([*s.notes, *page.extra_notes])


def _draw_header(slide, page: _Page) -> float:
    s = page.src
    heights = _header_heights(page.title, s.kicker if page.first else "", s.warning, s.by_rule)
    y = 0.25
    _label(slide, ML, y, CW, heights[0], page.title, _TITLE_SIZE, TX, bold=True, font=FH)
    slide.shapes[-1].name = "Title"
    y += heights[0]
    for height, text, size, color, bold, italic in (
        (heights[1], s.kicker if page.first else "", 12, TX_M, False, True),
        (heights[2], s.warning, 12, RED, True, False),
        (heights[3], _BY_RULE if s.by_rule else "", 10, TX_M, False, True),
    ):
        if text:
            _label(slide, ML, y, CW, height, text, size, color, bold=bold, italic=italic)
            y += height
    return y


def _draw_footer(slide, page: _Page, number: int, total: int) -> None:
    marks = {m.number: m for m in page.src.refs}
    items = _footer_marks(set(page.refs), marks)
    foot_h = _foot_h(items)
    y = SH - _FOOT_BOTTOM - foot_h
    _rect(slide, ML, y - 0.05, CW, 0.01, DIV).name = "Footer rule"
    if items:
        par = _frame(slide, ML, y, _FOOT_W, foot_h).paragraphs[0]
        slide.shapes[-1].name = "Footer references"
        _run(par, "References: ", 10, TX_M)
        for i, (n, mark) in enumerate(items):
            if i:
                _run(par, " · ", 10, TX_M)
            _run(par, f"[{n}] {mark.short}" if mark else f"[{n}]", 10, TX_M, link=mark.url if mark else None)
    _label(slide, SW - ML - 1.0, SH - _FOOT_BOTTOM - _FOOT_MIN, 1.0, _FOOT_MIN, f"{number} / {total}", 8, TX_M,
           align="right", anchor="middle")
    slide.shapes[-1].name = "Page number"


def _draw_title_page(slide, page: _Page) -> None:
    s = page.src
    _rect(slide, 0, 0, SW, 0.05, ACC)
    title_h = _need(s.title, 36, 11.3) + 0.1
    _label(slide, 1.0, 1.3, 11.3, title_h, s.title, 36, WH, bold=True, font=FH)
    y = 1.3 + title_h + 0.15
    _rect(slide, 1.0, y, 2.5, 0.03, ACC)
    y += 0.33
    for i, line in enumerate(s.subtitle.split("\n") if s.subtitle else []):
        size = 16 if i == 0 else 14
        h = _need(line, size, 11.3) + 0.08
        _label(slide, 1.0, y, 11.3, h, line, size, ACC_LT if i == 0 else BG)
        y += h
    if page.paras:
        y += 0.2
        tf = _frame(slide, 0.9, y, 11.5, page.text_h + 0.05, inset=0.1)
        _write_paras(tf, page.paras, 14, color=BG, muted=ACC_LT, bullet=ACC)
    if s.warning:
        h = _need(s.warning, 12, 11.1) + 0.2
        _rect(slide, 1.0, SH - 0.5 - h, 11.3, h, WH, rounded=True)
        _label(slide, 1.1, SH - 0.5 - h, 11.1, h, s.warning, 12, RED, bold=True, anchor="middle")


def _draw_flow(slide, page: _Page) -> None:
    s = page.src
    x = ML - 0.1
    y = page.top
    figure_below = 0.0
    if page.figure and page.variant in ("right", "left"):
        _draw_figure(slide, page, y)
        if page.variant == "left":
            x = ML + _IMG_COL + _IMG_GAP - 0.1
    if page.paras:
        tf = _frame(slide, x, y, page.text_w + 0.2, page.text_h + 0.05, inset=0.1)
        slide.shapes[-1].name = "Body"
        _write_paras(tf, page.paras, s.size, bullets=s.layout != "references")
        y += page.text_h + 0.15
    if page.rows:
        y = _draw_table(slide, page, y)
    if page.figure and page.variant == "top":
        _draw_figure(slide, page, y)
    elif page.visual:
        alone = not page.paras and not page.rows
        low, high = _visual_bounds(page.visual, alone)
        height = min(high, max(page.bottom - y, low))
        _draw_visual(slide, page.visual, y + 0.05, height)


def _draw_figure(slide, page: _Page, y: float) -> None:
    from pptx.util import Inches

    f, s = page.figure, page.src
    marks = {m.number: m for m in s.refs}
    cap, _ = _caption(f)
    col = CW if page.variant == "top" else _IMG_COL
    col_x = SW - ML - _IMG_COL if page.variant == "right" else ML
    below = _need(cap, 9, col) + 0.04 + _need(_credit_text(f, s.figure_ref, marks), 8, col) + 0.08
    pw, ph = page.fig_px
    avail_h = max(page.bottom - y - below - 0.05, 0.5)
    width = min(col, avail_h * pw / ph)
    height = width * ph / pw
    image_x = col_x + (col - width) / 2
    try:
        slide.shapes.add_picture(f.path, Inches(image_x), Inches(y), Inches(width), Inches(height)).name = "Figure"
    except Exception:  # noqa: BLE001 - an image python-pptx cannot embed is left out, the slide still draws
        return
    cy = y + height + 0.04
    _label(slide, col_x, cy, col, _need(cap, 9, col) + 0.04, cap, 9, TX_M, align="center")
    slide.shapes[-1].name = "Figure caption"
    cy += _need(cap, 9, col) + 0.04
    tf = _frame(slide, col_x, cy, col, _need(_credit_text(f, s.figure_ref, marks), 8, col) + 0.04)
    slide.shapes[-1].name = "Figure credit"
    par = tf.paragraphs[0]
    from pptx.enum.text import PP_ALIGN

    par.alignment = PP_ALIGN.CENTER
    mark = marks.get(s.figure_ref) if s.figure_ref is not None else None
    if s.figure_ref is not None:
        _run(par, f"{f.label}, ", 8, TX_M)
        _run(par, f"[{s.figure_ref}] {mark.short}" if mark else f"[{s.figure_ref}]", 8, TX_M, link=mark.url if mark else None)
    else:
        _run(par, f.label, 8, TX_M)
    if f.licence:
        _run(par, " - ", 8, TX_M)
        _run(par, f.licence, 8, TX_M, link=f.licence_url)


def _draw_table(slide, page: _Page, y: float) -> float:
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    s = page.src
    header = s.table[0]
    widths = _column_widths(s.table)
    rows = [header, *page.rows]
    heights = [_row_h(r, widths) for r in rows]
    shape = slide.shapes.add_table(len(rows), len(header), Inches(ML), Inches(y), Inches(CW), Inches(sum(heights)))
    table = shape.table
    table.first_row, table.horz_banding = True, False
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        table.rows[r].height = Inches(heights[r])
        fill = PRI if r == 0 else (CARD if r % 2 else ZEBRA)
        for c in range(len(header)):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(fill)
            par = cell.text_frame.paragraphs[0]
            _run(par, row[c] if c < len(row) else "", _TABLE_SIZE, _on(PRI) if r == 0 else TX, bold=r == 0)
    return y + sum(heights) + 0.15


# --- visual elements (ArcDeck's ve* functions) -----------------------------------------------------------------


def _fmt(value: float) -> str:
    return f"{value:g}"


def _draw_visual(slide, v: Visual, y: float, h: float) -> None:
    if v.caption:
        cap_h = _need(v.caption, 11, CW) + 0.1
        _label(slide, ML, y, CW, cap_h, v.caption, 11, TX_M, italic=True)
        y, h = y + cap_h, h - cap_h
    {
        "stats_callout": _ve_stats, "callout_cards": _ve_cards, "comparison": _ve_comparison,
        "stacked_list": _ve_stacked, "bar_chart": _ve_bar, "doughnut_chart": _ve_doughnut,
    }[v.type](slide, v.items, y, h)


def _card(slide, x, y, w, h, color):
    _rect(slide, x, y, w, h, CARD, rounded=True, line=DIV, shadow=True)
    _rect(slide, x, y, 0.06, h, color)


def _ve_stats(slide, items, y, h):
    items = items[:5]
    n, gap = len(items), 0.25
    cw = (CW - (n - 1) * gap) / n
    ch = min(h - 0.05, 2.2)
    num_h = ch * 0.55
    for i, item in enumerate(items):
        x, color = ML + i * (cw + gap), VE_COLS[i % len(VE_COLS)]
        big, small = (_fmt(item.value), item.label) if item.value is not None else (item.label, item.sublabel)
        if item.value is not None and item.sublabel:
            small = f"{item.label}\n{item.sublabel}"
        _rect(slide, x, y, cw, ch, CARD, rounded=True, line=DIV, shadow=True)
        _rect(slide, x, y, cw, 0.05, color)
        _label(slide, x + 0.1, y + 0.08, cw - 0.2, num_h, big, 34, color, bold=True, font=FH, align="center", anchor="middle")
        _label(slide, x + 0.1, y + 0.08 + num_h, cw - 0.2, ch - num_h - 0.16, small, 11, TX_M, align="center")


def _ve_cards(slide, items, y, h):
    items = items[:6]
    n, gap = len(items), 0.25
    cw = (CW - (n - 1) * gap) / n
    ch = min(h - 0.05, 2.8)
    for i, item in enumerate(items):
        x, color = ML + i * (cw + gap), VE_COLS[i % len(VE_COLS)]
        _card(slide, x, y, cw, ch, color)
        label_h = _need(item.label, 12, cw - 0.35)
        _label(slide, x + 0.2, y + 0.12, cw - 0.35, label_h, item.label, 12, DEEP, bold=True)
        if item.sublabel:
            _label(slide, x + 0.2, y + 0.2 + label_h, cw - 0.35, ch - label_h - 0.3, item.sublabel, 11, TX_M)


def _ve_comparison(slide, items, y, h):
    from pptx.enum.text import PP_ALIGN

    items = items[:5]
    n, gap = len(items), 0.2
    cw = (CW - (n - 1) * gap) / n
    head = max([0.45] + [_need(i.label, 13, cw - 0.2) + 0.14 for i in items])
    for i, item in enumerate(items):
        x = ML + i * (cw + gap)
        fill = TONE_COLS.get(item.tone or "neutral", PRI)
        _rect(slide, x, y, cw, h, CARD, rounded=True, line=DIV, shadow=True)
        _rect(slide, x, y, cw, head, fill)
        _label(slide, x + 0.1, y, cw - 0.2, head, item.label, 13, _on(fill), bold=True, align="center", anchor="middle")
        top = y + head + 0.1
        if item.value is not None:
            _label(slide, x + 0.1, top, cw - 0.2, 0.55, _fmt(item.value), 28, DEEP, bold=True, font=FH, align="center", anchor="middle")
            top += 0.55
        if item.sublabel:
            tf = _frame(slide, x + 0.1, top, cw - 0.2, y + h - top - 0.05)
            for j, line in enumerate(item.sublabel.split("\n")):
                par = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                par.alignment = PP_ALIGN.CENTER
                _run(par, line, 11, TX_M)


def _ve_stacked(slide, items, y, h):
    rows = _stack_rows(items)
    cy = y
    for i, (item, rh) in enumerate(zip(items[:8], rows)):
        color = VE_COLS[i % len(VE_COLS)]
        _card(slide, ML, cy, CW, rh, color)
        _label(slide, ML + 0.22, cy, _STACK_LABEL_W, rh, item.label, 13, DEEP, bold=True, anchor="middle")
        if item.sublabel:
            _label(slide, ML + 3.8, cy, _STACK_SUB_W, rh, item.sublabel, 12, TX_M, anchor="middle")
        cy += rh + 0.07


def _bar_color(item, i: int) -> str:
    """An outcome keeps its colour by meaning: the item's sublabel carries the
    outcome's own key ("no_effect") when the label is display text ("No effect")."""
    for text in (item.sublabel, item.label):
        key = text.strip().lower().replace(" ", "_").replace("-", "_")
        if key in OUTCOME_COLS:
            return OUTCOME_COLS[key]
    return TONE_COLS[item.tone] if item.tone else CHART_COLS[i % len(CHART_COLS)]


def _chart_data(items, labels):
    from pptx.chart.data import CategoryChartData

    data = CategoryChartData()
    data.categories = labels
    data.add_series("Count", [i.value for i in items])
    return data


def _ve_bar(slide, items, y, h):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Inches, Pt

    items = [i for i in items if i.value is not None][:8]
    _rect(slide, ML, y, CW, h, CARD, rounded=True, line=DIV)
    if not items:
        return
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(ML + 0.3), Inches(y + 0.05), Inches(CW - 0.6), Inches(h - 0.1), _chart_data(items, [i.label for i in items]))
    chart = frame.chart
    chart.has_legend = chart.has_title = False
    chart.font.name, chart.font.size = FB, Pt(10)
    chart.font.color.rgb = _rgb(TX_M)
    plot = chart.plots[0]
    plot.gap_width, plot.vary_by_categories = 80, False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position, labels.number_format, labels.number_format_is_linked = XL_LABEL_POSITION.OUTSIDE_END, "0", False
    labels.font.size, labels.font.color.rgb = Pt(10), _rgb(TX_M)
    for i, item in enumerate(items):
        point = plot.series[0].points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(_bar_color(item, i))
    axis = chart.value_axis
    axis.visible = False
    axis.minimum_scale, axis.maximum_scale = 0, max(1.0, max(i.value for i in items) * 1.3)
    axis.has_major_gridlines = True
    axis.major_gridlines.format.line.color.rgb = _rgb(DIV)
    chart.category_axis.format.line.color.rgb = _rgb(DIV)
    chart.category_axis.tick_labels.font.size = Pt(10)


def _ve_doughnut(slide, items, y, h):
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.oxml.ns import qn
    from pptx.util import Inches, Pt

    items = [i for i in items if i.value is not None][:8]
    _rect(slide, ML, y, CW, h, CARD, rounded=True, line=DIV)
    if not items:
        return
    w = min(CW - 0.6, max(h * 2.6, 6.0))
    labels = [f"{i.label} ({_fmt(i.value)})" for i in items]  # counts go in the legend: no number sits on a slice
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(ML + (CW - w) / 2), Inches(y + 0.05), Inches(w), Inches(h - 0.1), _chart_data(items, labels))
    chart = frame.chart
    chart.has_title = False
    chart.has_legend, chart.legend.position, chart.legend.include_in_layout = True, XL_LEGEND_POSITION.RIGHT, False
    chart.legend.font.size, chart.legend.font.name = Pt(11), FB
    chart.legend.font.color.rgb = _rgb(TX_M)
    plot = chart.plots[0]
    plot.vary_by_categories = True
    for i, item in enumerate(items):
        point = plot.series[0].points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = _rgb(_bar_color(item, i))
        point.format.line.color.rgb = _rgb(WH)
    hole = plot._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "55")


# --- the matrix (ArcDeck veMatrix, as Propeller's SWOT) ----------------------------------------------------------


def _draw_matrix(slide, page: _Page) -> None:
    size = min(page.src.size, 12)
    card_w = (CW - 0.25) / 2
    text_w = card_w - 0.3
    y = page.top
    for r, height in enumerate(_matrix_rows(page.groups, size, text_w, page.first)):
        if not height:
            continue
        for c, (label, items) in enumerate(page.groups[2 * r : 2 * r + 2]):
            if not items and not page.first:
                continue
            x, color = ML + c * (card_w + 0.25), SWOT_COLS[(2 * r + c) % len(SWOT_COLS)]
            _card(slide, x, y, card_w, height, color)
            _label(slide, x + 0.22, y + 0.1, card_w - 0.35, 0.3, label, 13, DEEP, bold=True, anchor="middle")
            tf = _frame(slide, x + 0.12, y + 0.5, text_w + 0.2, height - 0.5, inset=0.1)
            slide.shapes[-1].name = "Card text"
            _write_paras(tf, items or [Para("-", muted=True)], size, bullets=bool(items))
        y += height + 0.15
